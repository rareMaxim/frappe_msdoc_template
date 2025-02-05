from io import BytesIO
import frappe
import os
from docxtpl import DocxTemplate
from frappe.utils.file_manager import get_file_path
from frappe_msdoc_template.frappe_msdoc_template.doctype.document_template.document_template import DocumentTemplate
from openpyxl import load_workbook
from jinja2 import Template


class DocumentGenerator:
    def __init__(self, template_name, doctype, docname):
        self.template_name = template_name
        self.doctype = doctype
        self.docname = docname

    def generate(self):
        """Генерує документ на основі шаблону та DocType"""
        templates_info = get_templates(self.doctype)
        for template_info in templates_info:
            if template_info["name"] == self.template_name:
                frappe.logger().info(template_info)
                template_file = template_info["template"]
                # Отримуємо дані з DocType
                doc = frappe.get_doc(self.doctype, self.docname)
                self.context = doc.as_dict()
                if template_file.endswith(".docx"):
                    return self.generate_word(template_info)
                elif template_file.endswith(".xlsx"):
                    return self.generate_excel(template_info)
                else:
                    frappe.throw("Unsupported template format")

    def generate_word(self, template_info):
        """Генерує Word-документ на основі шаблону"""
        file_path = get_file_path(template_info["template"])
        doc = DocxTemplate(file_path)
        doc.render(self.context)
        output_stream = BytesIO()
        doc.save(output_stream)
        frappe.response["type"] = "binary"
        frappe.response["filecontent"] = output_stream.getvalue()
        frappe.response["filename"] = f"{template_info['title']} - {self.docname}.docx"
        frappe.response["type"] = "download"
        frappe.response["display_content_as"] = "attachment"
        output_stream.close()

    def generate_excel(self, template_info):
        """Генерує Excel-файл на основі шаблону"""
        file_path = get_file_path(template_info["template"])
        wb = load_workbook(file_path)
        ws = wb.active

        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value:
                    template = Template(cell.value)
                    cell.value = template.render(self.context)
        out_file = f"{template_info['title']} - {self.docname}.xlsx"
        wb.save(out_file)
        frappe.response["type"] = "binary"
        with open(out_file, "rb") as fh:
            frappe.response["filecontent"] = fh.read()
        frappe.response["filename"] = f"{template_info['title']} - {self.docname}.xlsx"
        frappe.response["type"] = "download"
        frappe.response["display_content_as"] = "attachment"


@frappe.whitelist()
def generate_document(template_name, doctype, docname):
    """Frappe API для генерації документів"""
    generator = DocumentGenerator(template_name, doctype, docname)
    return generator.generate()


@frappe.whitelist()
def get_link(template_name, doctype, docname):
    """Повертає посилання на документ"""
    return f"/api/method/frappe_msdoc_template.tmpl_gen.generate_document?template_name={template_name}&doctype={doctype}&docname={docname}"


@frappe.whitelist()
def get_templates(doctype):
    """Повертає список шаблонів для DocType"""
    doc: DocumentTemplate = frappe.get_last_doc(
        "Document Template", {"doc_type": doctype})
    result = []
    if not doc:
        return []
    result = []
    for d in doc.templates:
        result.append({
            "title": d.title,
            "name": d.name,
            "template": d.template,
            "group": d.group,
        })
    return result


@frappe.whitelist()
def get_doc_context(doctype, docname):
    """Повертає контекст документа"""
    doc = frappe.get_doc(doctype, docname)
    return doc.as_dict()
